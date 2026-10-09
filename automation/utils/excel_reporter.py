import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from typing import List, Dict, Any
from automation.config.config import Config

class ExcelReporter:
    @staticmethod
    def generate_excel_reports(results: List[Dict[str, Any]], output_dir: str = None) -> Dict[str, str]:
        target_dir = output_dir or os.path.join(Config.REPORTS_DIR, "Excel")
        os.makedirs(target_dir, exist_ok=True)

        passed_tests = [r for r in results if r["status"].upper() == "PASSED"]
        failed_tests = [r for r in results if r["status"].upper() == "FAILED"]
        skipped_tests = [r for r in results if r["status"].upper() in ["SKIPPED", "BLOCKED"]]

        # File 1: Automation_Test_Report.xlsx (6 Sheets)
        file1 = os.path.join(target_dir, "Automation_Test_Report.xlsx")
        wb1 = openpyxl.Workbook()

        # Styles
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )
        pass_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
        fail_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")

        # Sheet 1: Executed Test Cases
        ws1 = wb1.active
        ws1.title = "Executed Test Cases"
        headers1 = ["Test ID", "Module", "Test Name", "Status", "Execution Time", "Priority"]
        ws1.append(headers1)

        for row_idx, r in enumerate(results, start=2):
            ws1.append([
                r.get("test_id", ""),
                r.get("module", ""),
                r.get("name", ""),
                r.get("status", ""),
                f"{r.get('duration_ms', 0):.2f} ms",
                r.get("priority", "Medium")
            ])
            status_cell = ws1.cell(row=row_idx, column=4)
            if r.get("status", "").upper() == "PASSED":
                status_cell.fill = pass_fill
            elif r.get("status", "").upper() == "FAILED":
                status_cell.fill = fail_fill

        # Sheet 2: Passed Tests
        ws2 = wb1.create_sheet(title="Passed Tests")
        ws2.append(["Test ID", "Module", "Test Name", "Execution Time", "Priority"])
        for r in passed_tests:
            ws2.append([r["test_id"], r["module"], r["name"], f"{r.get('duration_ms', 0):.2f} ms", r.get("priority", "Medium")])

        # Sheet 3: Failed Tests
        ws3 = wb1.create_sheet(title="Failed Tests")
        ws3.append(["Test ID", "Module", "Test Name", "Failure Reason", "Screenshot Path"])
        for r in failed_tests:
            ws3.append([r["test_id"], r["module"], r["name"], r.get("failure_reason", "Assertion Failed"), r.get("screenshot", "")])

        # Sheet 4: Skipped Tests
        ws4 = wb1.create_sheet(title="Skipped Tests")
        ws4.append(["Test ID", "Module", "Test Name", "Skip Reason"])
        for r in skipped_tests:
            ws4.append([r["test_id"], r["module"], r["name"], r.get("failure_reason", "Precondition not met")])

        # Sheet 5: Execution Metrics
        ws5 = wb1.create_sheet(title="Execution Metrics")
        ws5.append(["Metric Name", "Value"])
        total = len(results)
        pass_rate = (len(passed_tests) / total * 100.0) if total > 0 else 0.0
        ws5.append(["Total Executed Test Cases", total])
        ws5.append(["Passed Test Cases", len(passed_tests)])
        ws5.append(["Failed Test Cases", len(failed_tests)])
        ws5.append(["Skipped Test Cases", len(skipped_tests)])
        ws5.append(["Pass Percentage", f"{pass_rate:.2f}%"])

        # Sheet 6: Defect Summary
        ws6 = wb1.create_sheet(title="Defect Summary")
        ws6.append(["Defect ID", "Test ID", "Module", "Severity", "Summary"])
        for idx, r in enumerate(failed_tests, start=1):
            ws6.append([f"DEF-{idx:03d}", r["test_id"], r["module"], "High", r.get("failure_reason", "Validation Error")])

        # Apply formatting across all sheets
        for sheet in wb1.worksheets:
            for cell in sheet[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    cell.border = thin_border

        wb1.save(file1)

        # File 2: Failed_Test_Cases.xlsx
        file2 = os.path.join(target_dir, "Failed_Test_Cases.xlsx")
        wb2 = openpyxl.Workbook()
        ws_f = wb2.active
        ws_f.title = "Failed Tests"
        ws_f.append(["Test ID", "Module", "Test Name", "Failure Reason", "Screenshot"])
        for r in failed_tests:
            ws_f.append([r["test_id"], r["module"], r["name"], r.get("failure_reason", ""), r.get("screenshot", "")])
        wb2.save(file2)

        # File 3: Passed_Test_Cases.xlsx
        file3 = os.path.join(target_dir, "Passed_Test_Cases.xlsx")
        wb3 = openpyxl.Workbook()
        ws_p = wb3.active
        ws_p.title = "Passed Tests"
        ws_p.append(["Test ID", "Module", "Test Name", "Duration MS"])
        for r in passed_tests:
            ws_p.append([r["test_id"], r["module"], r["name"], r.get("duration_ms", 0)])
        wb3.save(file3)

        # File 4: Summary_Report.xlsx
        file4 = os.path.join(target_dir, "Summary_Report.xlsx")
        wb4 = openpyxl.Workbook()
        ws_s = wb4.active
        ws_s.title = "Executive Summary"
        ws_s.append(["Metric", "Count"])
        ws_s.append(["Total Test Cases", len(results)])
        ws_s.append(["Passed", len(passed_tests)])
        ws_s.append(["Failed", len(failed_tests)])
        ws_s.append(["Skipped", len(skipped_tests)])
        ws_s.append(["Pass Rate %", f"{pass_rate:.2f}%"])
        wb4.save(file4)

        return {
            "Automation_Test_Report": file1,
            "Failed_Test_Cases": file2,
            "Passed_Test_Cases": file3,
            "Summary_Report": file4
        }
