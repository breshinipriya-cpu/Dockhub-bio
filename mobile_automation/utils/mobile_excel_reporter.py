import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from typing import List, Dict, Any
from mobile_automation.config.appium_config import AppiumConfig

class MobileExcelReporter:
    @staticmethod
    def generate_excel_reports(results: List[Dict[str, Any]], output_dir: str = None) -> Dict[str, str]:
        target_dir = output_dir or os.path.join(AppiumConfig.REPORTS_DIR, "Excel")
        os.makedirs(target_dir, exist_ok=True)

        passed_tests = [r for r in results if r["status"].upper() == "PASSED"]
        failed_tests = [r for r in results if r["status"].upper() == "FAILED"]
        skipped_tests = [r for r in results if r["status"].upper() in ["SKIPPED", "BLOCKED"]]

        total = len(results)
        pass_rate = (len(passed_tests) / total * 100.0) if total > 0 else 0.0

        # Styles
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )
        pass_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
        fail_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")

        # File 1: Automation_Test_Report.xlsx (7 Sheets)
        file1 = os.path.join(target_dir, "Automation_Test_Report.xlsx")
        wb1 = openpyxl.Workbook()

        # Sheet 1: Executed Test Cases
        ws1 = wb1.active
        ws1.title = "Executed Test Cases"
        ws1.append(["Test ID", "Module", "Test Name", "Priority", "Status", "Execution Time"])
        for idx, r in enumerate(results, start=2):
            ws1.append([
                r["test_id"],
                r["module"],
                r["name"],
                r.get("priority", "Medium"),
                r["status"],
                f"{r.get('duration_ms', 0):.2f} ms"
            ])
            status_cell = ws1.cell(row=idx, column=5)
            if r["status"].upper() == "PASSED":
                status_cell.fill = pass_fill
            elif r["status"].upper() == "FAILED":
                status_cell.fill = fail_fill

        # Sheet 2: Passed Tests
        ws2 = wb1.create_sheet(title="Passed Tests")
        ws2.append(["Test ID", "Module", "Test Name", "Execution Time", "Priority"])
        for r in passed_tests:
            ws2.append([r["test_id"], r["module"], r["name"], f"{r.get('duration_ms', 0):.2f} ms", r.get("priority", "Medium")])

        # Sheet 3: Failed Tests
        ws3 = wb1.create_sheet(title="Failed Tests")
        ws3.append(["Test ID", "Module", "Test Name", "Failure Reason", "Screenshot Path"])
        for r in failed_tests:
            ws3.append([r["test_id"], r["module"], r["name"], r.get("failure_reason", "Assertion Mismatch"), r.get("screenshot", "")])

        # Sheet 4: Skipped Tests
        ws4 = wb1.create_sheet(title="Skipped Tests")
        ws4.append(["Test ID", "Module", "Test Name", "Skip Reason"])
        for r in skipped_tests:
            ws4.append([r["test_id"], r["module"], r["name"], r.get("failure_reason", "Feature Disabled")])

        # Sheet 5: Execution Metrics
        ws5 = wb1.create_sheet(title="Execution Metrics")
        ws5.append(["Metric Name", "Value"])
        ws5.append(["Total Test Cases", total])
        ws5.append(["Passed Test Cases", len(passed_tests)])
        ws5.append(["Failed Test Cases", len(failed_tests)])
        ws5.append(["Skipped Test Cases", len(skipped_tests)])
        ws5.append(["Pass Rate Percentage", f"{pass_rate:.2f}%"])

        # Sheet 6: Defect Summary
        ws6 = wb1.create_sheet(title="Defect Summary")
        ws6.append(["Defect ID", "Test ID", "Module", "Severity", "Summary"])
        for idx, r in enumerate(failed_tests, start=1):
            ws6.append([f"DEF-MOB-{idx:03d}", r["test_id"], r["module"], "High", r.get("failure_reason", "Appium Assertion Mismatch")])

        # Sheet 7: Pass Rate Summary
        ws7 = wb1.create_sheet(title="Pass Rate Summary")
        ws7.append(["Category", "Total", "Passed", "Failed", "Pass Rate %"])
        modules = {}
        for r in results:
            m = r["module"]
            if m not in modules:
                modules[m] = {"total": 0, "passed": 0, "failed": 0}
            modules[m]["total"] += 1
            if r["status"].upper() == "PASSED":
                modules[m]["passed"] += 1
            else:
                modules[m]["failed"] += 1
        for m, c in modules.items():
            rate = (c["passed"] / c["total"] * 100.0) if c["total"] > 0 else 0.0
            ws7.append([m, c["total"], c["passed"], c["failed"], f"{rate:.1f}%"])

        # Format Header Styles
        for sheet in wb1.worksheets:
            for cell in sheet[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    cell.border = thin_border

        wb1.save(file1)

        # File 2: Passed_Test_Cases.xlsx
        file2 = os.path.join(target_dir, "Passed_Test_Cases.xlsx")
        wb2 = openpyxl.Workbook()
        ws_p = wb2.active
        ws_p.title = "Passed Tests"
        ws_p.append(["Test ID", "Module", "Test Name", "Duration MS"])
        for r in passed_tests:
            ws_p.append([r["test_id"], r["module"], r["name"], r.get("duration_ms", 0)])
        wb2.save(file2)

        # File 3: Failed_Test_Cases.xlsx
        file3 = os.path.join(target_dir, "Failed_Test_Cases.xlsx")
        wb3 = openpyxl.Workbook()
        ws_f = wb3.active
        ws_f.title = "Failed Tests"
        ws_f.append(["Test ID", "Module", "Test Name", "Failure Reason", "Screenshot"])
        for r in failed_tests:
            ws_f.append([r["test_id"], r["module"], r["name"], r.get("failure_reason", ""), r.get("screenshot", "")])
        wb3.save(file3)

        # File 4: Execution_Summary.xlsx
        file4 = os.path.join(target_dir, "Execution_Summary.xlsx")
        wb4 = openpyxl.Workbook()
        ws_s = wb4.active
        ws_s.title = "Summary"
        ws_s.append(["Metric", "Value"])
        ws_s.append(["Total Executed", total])
        ws_s.append(["Passed", len(passed_tests)])
        ws_s.append(["Failed", len(failed_tests)])
        ws_s.append(["Skipped", len(skipped_tests)])
        ws_s.append(["Pass Percentage", f"{pass_rate:.2f}%"])
        wb4.save(file4)

        return {
            "Automation_Test_Report": file1,
            "Passed_Test_Cases": file2,
            "Failed_Test_Cases": file3,
            "Execution_Summary": file4
        }
